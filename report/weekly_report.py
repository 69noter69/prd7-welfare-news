# -*- coding: utf-8 -*-
"""สร้างรายงานประจำสัปดาห์เป็นไฟล์ Excel แล้วส่งเข้าอีเมล

รันอัตโนมัติทุกเช้าวันจันทร์ 08:30 น. โดย GitHub Actions
ช่วงข้อมูล: จันทร์ - อาทิตย์ ของสัปดาห์ที่เพิ่งจบ

    python -m report.weekly_report              # สร้างและส่งอีเมล
    python -m report.weekly_report --dry-run    # สร้างอย่างเดียว ไม่ส่ง (ไว้ทดสอบ)
    python -m report.weekly_report --week 2026-09-21   # ระบุวันจันทร์ต้นสัปดาห์เอง
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import smtplib
import ssl
import sys
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OUT = ROOT / "report" / "out"

FONT = "Tahoma"
INK, MUTED = "1A1D23", "5B6472"
BLUE, ORANGE = "2563A8", "C2410C"
LINE, BAND = "D6DCE5", "EEF3F9"
GOOD = "1A6F47"

THAI_DAY = ["จ.", "อ.", "พ.", "พฤ.", "ศ.", "ส.", "อา."]
THAI_MONTH = ["", "ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.",
              "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]


def th_short(d: date) -> str:
    return f"{d.day} {THAI_MONTH[d.month]} {d.year + 543}"


def th_day_label(d: date) -> str:
    return f"{THAI_DAY[d.weekday()]} {d.day} {THAI_MONTH[d.month]}"


# ---------------------------------------------------------------- อ่านข้อมูล
def read_rows() -> list[dict]:
    """อ่าน data/stats-daily.csv ที่ระบบเก็บข่าวเขียนไว้"""
    path = DATA / "stats-daily.csv"
    if not path.exists():
        sys.exit(f"ไม่พบไฟล์ {path} — ระบบเก็บข่าวยังไม่เคยรันสำเร็จ")
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def num(value, default=0):
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return default


def as_row(raw: dict) -> dict | None:
    try:
        d = datetime.strptime(raw["วันที่"].strip(), "%Y-%m-%d").date()
    except (KeyError, ValueError):
        return None
    return {
        "date": d,
        "runs": int(num(raw.get("ครั้งที่ค้นหา"), 1)),
        "fetched": int(num(raw.get("ข่าวดิบที่อ่าน"))),
        "kept": int(num(raw.get("ข่าวที่คัดมาแสดง"))),
        "new": int(num(raw.get("ข่าวใหม่"))),
        "followups": int(num(raw.get("ข่าวตามต่อ"))),
        "eastern": int(num(raw.get("ภูมิภาค"))),
        "national": int(num(raw.get("ทั่วประเทศ"))),
        "sources": int(num(raw.get("จำนวนแหล่งข่าว"))),
        "cost": round(num(raw.get("ค่าใช้จ่ายโดยประมาณ (บาท)")), 2),
    }


def last_full_week(today: date) -> date:
    """คืนวันจันทร์ของสัปดาห์ที่เพิ่งจบ (จันทร์-อาทิตย์)"""
    this_monday = today - timedelta(days=today.weekday())
    return this_monday - timedelta(days=7)


def slice_week(rows: list[dict], monday: date) -> list[dict]:
    sunday = monday + timedelta(days=6)
    return sorted((r for r in rows if monday <= r["date"] <= sunday), key=lambda r: r["date"])


def totals(week: list[dict]) -> dict:
    keys = ("runs", "fetched", "kept", "new", "followups", "eastern", "national", "sources")
    out = {k: sum(r[k] for r in week) for k in keys}
    out["cost"] = round(sum(r["cost"] for r in week), 2)
    out["days"] = len(week)
    return out


def read_index() -> dict:
    path = DATA / "index.json"
    if not path.exists():
        return {}
    try:
        with open(path, encoding="utf-8") as fh:
            return (json.load(fh) or {}).get("totals") or {}
    except (json.JSONDecodeError, OSError):
        return {}


# ---------------------------------------------------------------- สร้าง Excel
def build_xlsx(week: list[dict], prev: dict, idx: dict, monday: date, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    from openpyxl.comments import Comment

    def F(sz=10, b=False, color=INK, it=False):
        return Font(name=FONT, size=sz, bold=b, color=color, italic=it)
    thin = Side(style="thin", color=LINE)
    sunday = monday + timedelta(days=6)
    wb = Workbook()

    # ---------- ชีต 2 : ที่มาและวิธีใช้ ----------
    src = wb.create_sheet("ที่มาและวิธีใช้")
    src.sheet_view.showGridLines = False
    src["A1"] = "ที่มาของข้อมูล · ข้อสมมติ · ยอดสัปดาห์ก่อน"
    src["A1"].font = F(13, True)
    r = 3
    for a, b in [("แหล่งข้อมูล", "data/stats-daily.csv (ระบบเขียนเองทุกวัน)"),
                 ("และ", "data/index.json"),
                 ("สร้างรายงานเมื่อ", th_short(date.today())),
                 ("ช่วงข้อมูล", f"{th_short(monday)} – {th_short(sunday)}"),
                 ("ผู้จัดทำระบบ", "69noter69 รติ สปส สปท7")]:
        src.cell(row=r, column=1, value=a).font = F(10, True)
        src.cell(row=r, column=2, value=b).font = F(10, color=MUTED)
        r += 1

    r += 1
    src.cell(row=r, column=1, value="ข้อสมมติ — แก้ตัวเลขในช่องสีเหลืองได้ รายงานคำนวณใหม่เอง").font = F(11, True)
    r += 1

    def yellow(row, label, val, note=None):
        src.cell(row=row, column=1, value=label).font = F(10)
        c = src.cell(row=row, column=2, value=val)
        c.font = F(11, True, "0000FF")
        c.fill = PatternFill("solid", fgColor="FFFF00")
        c.alignment = Alignment(horizontal="center")
        if note:
            c.comment = Comment(note, "ระบบ")
        return f"'ที่มาและวิธีใช้'!$B${row}"

    SEC = yellow(r, "เวลาที่คนใช้อ่านและตัดสินข่าว 1 ชิ้น (วินาที)", 20,
                 "เป็นการประมาณการ ไม่ได้จับเวลาจริง\nถ้าอยากให้แม่น ให้จับเวลาจริงสัก 20 ข่าวแล้วแก้ตัวเลขนี้")
    r += 2
    src.cell(row=r, column=1, value="ยอดสัปดาห์ก่อน — ใช้คำนวณช่องเปรียบเทียบ").font = F(11, True)
    r += 1
    P = {}
    for key, label in [("fetched", "ข่าวดิบที่อ่าน"), ("kept", "ข่าวที่คัดมาแสดง"),
                       ("eastern", "ข่าวภูมิภาค"), ("cost", "ต้นทุนรวม (บาท)")]:
        src.cell(row=r, column=1, value=label).font = F(10)
        c = src.cell(row=r, column=2, value=prev.get(key, 0))
        c.font = F(10, color=MUTED)
        c.alignment = Alignment(horizontal="center")
        c.number_format = '#,##0.00' if key == "cost" else '#,##0'
        P[key] = f"'ที่มาและวิธีใช้'!$B${r}"
        r += 1

    r += 1
    for t, bold in [("ตัวเลขนี้มาจากไหน", True),
                    ("ระบบเก็บข่าวเขียนไฟล์สถิติเองทุกครั้งที่ทำงาน รายงานนี้แค่หยิบมาจัดรูปแบบ", False),
                    ("ไม่มีการกรอกมือ จึงไม่มีโอกาสพิมพ์ผิด", False),
                    ("", False),
                    ("ความลับของข้อมูล", True),
                    ("ไฟล์นี้ส่งเข้าอีเมลโดยตรง ไม่ได้เก็บไว้ใน repo ซึ่งเป็นแบบ Public", False)]:
        src.cell(row=r, column=1, value=t).font = F(10, bold)
        r += 1
    src.column_dimensions["A"].width = 56
    src.column_dimensions["B"].width = 52

    # ---------- ชีต 1 : รายงานสรุป ----------
    ws = wb.active
    ws.title = "รายงานสรุป"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJ", [15, 10, 12, 12, 10, 10, 10, 11, 12, 12]):
        ws.column_dimensions[col].width = w

    ws.merge_cells("A1:J1")
    ws["A1"] = "รายงานประจำสัปดาห์ · ระบบคัดกรองข่าวสวัสดิการอัตโนมัติ"
    ws["A1"].font = F(16, True, BLUE)
    ws.row_dimensions[1].height = 26
    ws.merge_cells("A2:J2")
    ws["A2"] = (f"สำนักงานประชาสัมพันธ์ที่ 7 · สัปดาห์ที่ {th_short(monday)} – {th_short(sunday)}"
                f" · ออกรายงาน {th_short(date.today())} เวลา 08:30 น.")
    ws["A2"].font = F(9.5, color=MUTED)
    ws.row_dimensions[3].height = 7

    # ตารางรายวัน
    TH = 13
    head = ["วันที่", "ครั้งที่ค้นหา", "ข่าวดิบที่อ่าน", "ข่าวที่คัดมาแสดง", "ข่าวใหม่",
            "ข่าวตามต่อ", "ภูมิภาค", "ทั่วประเทศ", "จำนวนแหล่งข่าว", "ต้นทุน (บาท)"]
    ws.merge_cells(f"A{TH-1}:J{TH-1}")
    ws[f"A{TH-1}"] = "รายละเอียดรายวัน"
    ws[f"A{TH-1}"].font = F(12, True)
    ws.row_dimensions[TH-1].height = 22
    for c, h in enumerate(head, 1):
        cell = ws.cell(row=TH, column=c, value=h)
        cell.font = F(9, True, "FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[TH].height = 30

    FR = TH + 1
    for i, row in enumerate(week):
        rr = FR + i
        weekend = row["date"].weekday() >= 5
        values = [th_day_label(row["date"]), row["runs"], row["fetched"], row["kept"],
                  row["new"], row["followups"], row["eastern"], row["national"],
                  row["sources"], row["cost"]]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=rr, column=c, value=v)
            cell.font = F(10, color=MUTED if weekend else INK)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(horizontal="left" if c == 1 else "center")
            cell.number_format = '#,##0.00' if c == 10 else ('#,##0' if c > 1 else 'General')
            if weekend:
                cell.fill = PatternFill("solid", fgColor="F7F8FA")
    LR = FR + len(week) - 1
    TR = LR + 1
    ws.cell(row=TR, column=1, value="รวมทั้งสัปดาห์").font = F(10, True)
    for c in range(2, 11):
        L = get_column_letter(c)
        cell = ws.cell(row=TR, column=c, value=f"=SUM({L}{FR}:{L}{LR})")
        cell.font = F(10, True)
        cell.alignment = Alignment(horizontal="center")
        cell.number_format = '#,##0.00' if c == 10 else '#,##0'
    for c in range(1, 11):
        cell = ws.cell(row=TR, column=c)
        cell.fill = PatternFill("solid", fgColor=BAND)
        cell.border = Border(top=Side(style="medium", color=BLUE))
    ws.merge_cells(f"A{TR+1}:J{TR+1}")
    ws[f"A{TR+1}"] = "หมายเหตุ: แถวสีเทาคือวันเสาร์–อาทิตย์ ซึ่งปริมาณข่าวน้อยกว่าวันทำงานตามปกติ"
    ws[f"A{TR+1}"].font = F(8.5, color=MUTED, it=True)

    # แถบตัวเลขสำคัญ
    def delta(now, prev_ref):
        return (f'=IF({prev_ref}=0,"เทียบไม่ได้ (ยังไม่มีข้อมูลสัปดาห์ก่อน)",'
                f'IF({now}>={prev_ref},"▲ ","▼ ")&TEXT(ABS({now}/{prev_ref}-1),"0.0%")&" จากสัปดาห์ก่อน")')
    tiles = [("A", "C", "ข่าวที่ส่งถึงทีม", f"=D{TR}", '#,##0" ข่าว"', delta(f"D{TR}", P["kept"])),
             ("D", "E", "ข่าวดิบที่ระบบอ่านแทนคน", f"=C{TR}", '#,##0" ข่าว"', delta(f"C{TR}", P["fetched"])),
             ("F", "G", "ต้นทุนรวมทั้งสัปดาห์", f"=J{TR}", '#,##0.00" บาท"', delta(f"J{TR}", P["cost"])),
             ("H", "J", "ต้นทุนต่อข่าว 1 ชิ้น", f'=IF(D{TR}=0,"",J{TR}/D{TR})', '0.00" บาท"',
              "ตัวเลขชี้วัดความคุ้มค่าหลัก")]
    for c1, c2, label, formula, fmt, sub in tiles:
        for rr in (5, 6, 7):
            ws.merge_cells(f"{c1}{rr}:{c2}{rr}")
        l = ws[f"{c1}5"]; l.value = label; l.font = F(9, color=MUTED)
        v = ws[f"{c1}6"]; v.value = formula; v.font = F(20, True, BLUE); v.number_format = fmt
        d = ws[f"{c1}7"]; d.value = sub; d.font = F(8.5, color=MUTED)
        for cell in (l, v, d):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for cc in range(ord(c1) - 64, ord(c2) - 63):
            for rr in (5, 6, 7):
                ws.cell(row=rr, column=cc).fill = PatternFill("solid", fgColor=BAND)
            ws.cell(row=7, column=cc).border = Border(bottom=Side(style="medium", color=BLUE))
    ws.row_dimensions[5].height = 16
    ws.row_dimensions[6].height = 30
    ws.row_dimensions[7].height = 15
    ws.row_dimensions[8].height = 8

    # สรุปส่งท้าย 2 บรรทัด
    ws.merge_cells("A9:J9")
    ws["A9"] = "สรุปส่งท้าย"
    ws["A9"].font = F(12, True)
    L1 = (f'=CONCATENATE("ระบบทำงาน ",TEXT(COUNT(C{FR}:C{LR}),"0")," วันในสัปดาห์นี้ '
          f'อ่านข่าว ",TEXT(C{TR},"#,##0")," ชิ้น คัดเหลือ ",TEXT(D{TR},"#,##0")," ชิ้น '
          f'ด้วยต้นทุน ",TEXT(J{TR},"#,##0.00")," บาท หรือ ",TEXT(J{TR}/D{TR}*100,"0")," สตางค์ต่อข่าว 1 ชิ้น '
          f'ประหยัดเวลาทีมงานราว ",TEXT((C{TR}-D{TR})*{SEC}/3600,"0")," ชั่วโมง")')
    L2 = (f'=CONCATENATE("ข่าวในพื้นที่ 9 จังหวัดที่รับผิดชอบมี ",TEXT(G{TR},"#,##0"),'
          f'" ชิ้น คิดเป็น ",TEXT(IF(D{TR}=0,0,G{TR}/D{TR}),"0.0%")," ของทั้งหมด",'
          f'IF(IF(D{TR}=0,0,G{TR}/D{TR})<0.1," ซึ่งยังต่ำกว่าที่ควรเป็นสำหรับสำนักงานที่ดูแลพื้นที่นี้โดยตรง",'
          f'" ถือว่าอยู่ในเกณฑ์ที่ใช้งานได้"))')
    for i, f in enumerate((L1, L2)):
        rr = 10 + i
        ws.merge_cells(f"A{rr}:J{rr}")
        cell = ws[f"A{rr}"]
        cell.value = f
        cell.font = F(10.5)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        for cc in range(1, 11):
            ws.cell(row=rr, column=cc).fill = PatternFill("solid", fgColor="F4F8FD")
        ws.row_dimensions[rr].height = 30 if i == 0 else 18
        ws[f"A{rr}"].border = Border(left=Side(style="thick", color=BLUE))

    # กราฟ
    ch = BarChart()
    ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100
    ch.title = "ข่าวที่คัดมาแสดงแต่ละวัน แยกตามพื้นที่"
    ch.y_axis.title = "จำนวนข่าว"
    ch.height = 6.8; ch.width = 17.8; ch.gapWidth = 55
    ch.add_data(Reference(ws, min_col=7, max_col=8, min_row=TH, max_row=LR), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=FR, max_row=LR))
    ch.series[0].graphicalProperties = GraphicalProperties(solidFill=ORANGE)
    ch.series[1].graphicalProperties = GraphicalProperties(solidFill=BLUE)
    for s in ch.series:
        s.graphicalProperties.line = LineProperties(solidFill="FFFFFF", w=25400)
    ch.series[0].dLbls = DataLabelList()
    ch.series[0].dLbls.showVal = True
    for a in ("showSerName", "showCatName", "showLegendKey"):
        setattr(ch.series[0].dLbls, a, False)
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINE, w=9525))
    ws.add_chart(ch, f"A{TR+3}")

    FT = TR + 21
    ws.merge_cells(f"A{FT}:J{FT}")
    ws[f"A{FT}"] = ("ที่มา: ไฟล์ stats-daily.csv และ index.json ที่ระบบเขียนเองอัตโนมัติ · "
                    "ตัวเลขเวลาที่ประหยัดได้คำนวณจากข้อสมมติในชีต 'ที่มาและวิธีใช้' ซึ่งแก้ได้เอง")
    ws[f"A{FT}"].font = F(8.5, color=MUTED, it=True)
    ws[f"A{FT}"].alignment = Alignment(wrap_text=True, vertical="top")

    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.move_sheet("รายงานสรุป", offset=-1)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------------------------------------------------------------- ส่งอีเมล
def email_body(cur: dict, prev: dict, monday: date, sunday: date, streak) -> str:
    def line(label, value, prev_value, fmt="{:,.0f}"):
        text = fmt.format(value)
        if prev_value:
            pct = (value / prev_value - 1) * 100
            text += f"  ({'▲' if pct >= 0 else '▼'} {abs(pct):.1f}% จากสัปดาห์ก่อน)"
        return f"  {label:<28}{text}"

    per_item = (cur["cost"] / cur["kept"]) if cur["kept"] else 0
    eastern_pct = (cur["eastern"] / cur["kept"] * 100) if cur["kept"] else 0
    saved_hours = (cur["fetched"] - cur["kept"]) * 20 / 3600

    return "\n".join([
        f"รายงานประจำสัปดาห์ {th_short(monday)} – {th_short(sunday)}",
        "สำนักงานประชาสัมพันธ์ที่ 7",
        "",
        "ตัวเลขสำคัญ",
        line("ข่าวที่ส่งถึงทีม", cur["kept"], prev.get("kept")) + " ข่าว",
        line("ข่าวดิบที่ระบบอ่านแทนคน", cur["fetched"], prev.get("fetched")) + " ข่าว",
        line("ต้นทุนรวม", cur["cost"], prev.get("cost"), "{:,.2f}") + " บาท",
        f"  {'ต้นทุนต่อข่าว 1 ชิ้น':<28}{per_item:.2f} บาท",
        "",
        f"  ข่าวในพื้นที่ 9 จังหวัด     {cur['eastern']:,} ข่าว ({eastern_pct:.1f}% ของทั้งหมด)",
        f"  ประหยัดเวลาทีมงาน          ประมาณ {saved_hours:.0f} ชั่วโมง",
        f"  ระบบทำงาน                  {cur['days']} วัน"
        + (f" · รันสำเร็จติดต่อกัน {streak} วัน" if streak else ""),
        "",
        "รายละเอียดทั้งหมดอยู่ในไฟล์ Excel ที่แนบมาด้วย",
        "",
        "— ส่งอัตโนมัติจากระบบเก็บข่าวสวัสดิการ ไม่ต้องตอบกลับ",
    ])


def send_mail(path: Path, subject: str, body: str) -> None:
    user = os.environ.get("MAIL_USERNAME", "").strip()
    password = os.environ.get("MAIL_PASSWORD", "").replace(" ", "").strip()
    to = os.environ.get("MAIL_TO", "").strip() or user
    if not user or not password:
        sys.exit("ไม่พบ MAIL_USERNAME หรือ MAIL_PASSWORD — ตั้งค่าใน Settings > Secrets ของ repo")

    msg = EmailMessage()
    msg["From"] = user
    msg["To"] = ", ".join(a.strip() for a in to.split(",") if a.strip())
    msg["Subject"] = subject
    msg.set_content(body)
    with open(path, "rb") as fh:
        msg.add_attachment(fh.read(), maintype="application",
                           subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           filename=path.name)
    context = ssl.create_default_context()
    with smtplib.SMTP("smtp.gmail.com", 587, timeout=60) as smtp:
        smtp.starttls(context=context)
        smtp.login(user, password)
        smtp.send_message(msg)
    print(f"ส่งอีเมลถึง {msg['To']} เรียบร้อย")


# ---------------------------------------------------------------- main
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="รายงานประจำสัปดาห์ สปท.7")
    ap.add_argument("--dry-run", action="store_true", help="สร้างไฟล์อย่างเดียว ไม่ส่งอีเมล")
    ap.add_argument("--week", help="วันจันทร์ต้นสัปดาห์ที่ต้องการ (YYYY-MM-DD)")
    args = ap.parse_args(argv)

    rows = [r for r in (as_row(x) for x in read_rows()) if r]
    if not rows:
        sys.exit("ไฟล์สถิติว่างเปล่า ยังไม่มีข้อมูลให้ทำรายงาน")

    monday = (datetime.strptime(args.week, "%Y-%m-%d").date() if args.week
              else last_full_week(date.today()))
    monday -= timedelta(days=monday.weekday())      # กันกรณีใส่วันที่ไม่ใช่วันจันทร์
    sunday = monday + timedelta(days=6)

    week = slice_week(rows, monday)
    if not week:
        print(f"ไม่มีข้อมูลในช่วง {monday} ถึง {sunday} — ข้ามรอบนี้")
        return 0

    prev = totals(slice_week(rows, monday - timedelta(days=7)))
    cur = totals(week)
    idx = read_index()

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / f"weekly-report-{monday}-to-{sunday}.xlsx"
    build_xlsx(week, prev, idx, monday, path)
    print(f"สร้างไฟล์ {path.name} ({cur['days']} วัน, ข่าว {cur['kept']} ชิ้น, {cur['cost']} บาท)")

    if args.dry_run:
        print("โหมดทดสอบ — ไม่ส่งอีเมล")
        return 0

    subject = f"[สปท.7] รายงานประจำสัปดาห์ {th_short(monday)} – {th_short(sunday)}"
    send_mail(path, subject, email_body(cur, prev, monday, sunday, idx.get("success_streak_days")))
    return 0


if __name__ == "__main__":
    sys.exit(main())
